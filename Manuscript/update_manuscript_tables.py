#!/usr/bin/env python3
"""Write core-method tables into both manuscript Word files and refresh Figure 1."""

import csv
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path("/Users/richardjackson/Documents/GitHub/synthetic-Controlled-Trials")
SA_CSV = ROOT / "Manuscript" / "Table_fully_synthetic_core.csv"
HY_CSV = ROOT / "Manuscript" / "Table_hybrid_core.csv"
FIG = ROOT / "Examples" / "psc_sim" / "results" / "coverage_vs_bias_sa_ra.png"

HEADER = ["Method", "Sample size", "Overlap", "Mean (se)", "ACIL", "Coverage", "Bias"]

SA_CAPTION = (
    "Table 2. Fully synthetic trials: mean estimate (SE), average interval length (ACIL), "
    "coverage of nominal 95% intervals (%), and bias for the log hazard ratio (true value log 0.7). "
    "Overlap Small/Medium/Large corresponds to none/moderate/large covariate overlap. "
    "Unadjusted Cox is a pooled Cox model. Core methods only (partial CFM, proximity-weighted PSC, "
    "and adjusted Cox models omitted)."
)
HY_CAPTION = (
    "Table 3. Partially synthetic (hybrid) trials: mean estimate (SE), ACIL, coverage (%), and bias "
    "for the log hazard ratio (true value log 0.7). Unadjusted Cox is the randomised-trial Cox model "
    "(RCT unadjusted); the pooled unadjusted Cox is not shown. Overlap Small/Medium/Large corresponds "
    "to none/moderate/large covariate overlap."
)
FIG_CAPTION = (
    "Figure 1. Coverage versus mean bias for core methods in fully synthetic (left) and "
    "partially synthetic / hybrid (right) scenarios. Methods are personalised "
    "synthetic controls, entropy-balanced synthetic controls, Bayesian case-weighted models, "
    "and (hybrid only) Bayesian commensurate priors. Each method appears as nine points "
    "per panel (three sample sizes × three overlap levels); commensurate priors are "
    "hybrid only. Panels use separate coverage scales so hybrid points are not compressed. "
    "The dashed horizontal line marks nominal 95% coverage; the dotted vertical line marks "
    "zero bias. Colour denotes method, shape sample size, and fill covariate overlap "
    "(open = small/none, translucent = moderate, solid = large). Points are offset "
    "slightly by sample size (horizontal) and overlap (vertical) so that all nine "
    "scenarios per method remain visible."
)
POINTER = (
    "Figure 1 plots coverage against bias for core methods in fully synthetic and hybrid settings "
    "across all eighteen scenarios (nine points per method in each panel: three sample sizes × "
    "three overlap levels). Table 2 reports fully synthetic results by sample size and overlap; "
    "Table 3 reports the corresponding hybrid results. Unadjusted Cox is a pooled Cox model in fully "
    "synthetic trials and the randomised-trial Cox model in hybrid trials."
)


def read_csv(path):
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def display_rows(records):
    """Blank repeated Method / Sample size to match nested template."""
    rows = []
    prev_m = prev_s = None
    for rec in records:
        m = rec["Method"]
        s = rec["Sample size"]
        rows.append([
            m if m != prev_m else "",
            s if not (m == prev_m and s == prev_s) else "",
            rec["Overlap"],
            rec["Mean (se)"],
            rec["ACIL"],
            rec["Coverage"],
            rec["Bias"],
        ])
        prev_m, prev_s = m, s
    return rows


def set_cell(cell, text):
    cell.text = text


def rebuild_table(table, header, data_rows):
    data = [header] + data_rows
    ncols = len(header)
    while len(table.rows) > len(data):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(data):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.columns) < ncols:
        tbl = table._tbl
        tbl.tblGrid.append(deepcopy(tbl.tblGrid.gridCol_lst[-1]))
        for tr in tbl.tr_lst:
            tr.append(deepcopy(tr.tc_lst[-1]))
    # drop extra columns
    while len(table.columns) > ncols:
        tbl = table._tbl
        grid = tbl.tblGrid
        grid.remove(grid.gridCol_lst[-1])
        for tr in tbl.tr_lst:
            tr.remove(tr.tc_lst[-1])
    for i, rowvals in enumerate(data):
        for j, val in enumerate(rowvals):
            set_cell(table.rows[i].cells[j], val)


def set_para(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.text = text


def replace_images(doc, png_path):
    blob = png_path.read_bytes()
    n = 0
    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        # Replace the large coverage-vs-bias figures, not the small schematic.
        if len(part.blob) > 100000:
            part._blob = blob
            n += 1
    return n


def update_doc(path, sa_rows, hy_rows):
    doc = Document(str(path))
    tables = []
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        tables.append((t, hdr))

    sa_table = hy_table = None
    for t, hdr in tables:
        joined = " ".join(hdr)
        if "Overlap" in hdr and "Sample size" in hdr and "Bias" in joined:
            if sa_table is None:
                sa_table = t
            elif hy_table is None and t is not sa_table:
                # second scenario table currently Table 3 method ranges
                hy_table = t
        elif hdr[:4] == ["Trial type", "Method", "Coverage (%)", "Bias"]:
            hy_table = t

    if sa_table is None or hy_table is None:
        raise SystemExit(f"Could not find tables in {path.name}: sa={sa_table is not None} hy={hy_table is not None}")

    rebuild_table(sa_table, HEADER, sa_rows)
    rebuild_table(hy_table, HEADER, hy_rows)

    for p in doc.paragraphs:
        t = p.text
        if t.startswith("Table 2."):
            set_para(p, SA_CAPTION)
        elif t.startswith("Table 3."):
            set_para(p, HY_CAPTION)
        elif t.startswith("Figure 1."):
            set_para(p, FIG_CAPTION)
        elif t.startswith("Figure 1 plots coverage"):
            set_para(p, POINTER)

    nimg = replace_images(doc, FIG)
    doc.save(str(path))
    print(f"{path.name}: SA rows {len(sa_table.rows)} HY rows {len(hy_table.rows)} images {nimg}")


def main():
    sa_rows = display_rows(read_csv(SA_CSV))
    hy_rows = display_rows(read_csv(HY_CSV))
    files = [
        ROOT / "Manuscript" / "Synthetically Controlled Trials in Cancer Research cursor_sugg.docx",
        ROOT / "Manuscript" / "Synthetically Controlled Trials august 26.docx",
    ]
    for path in files:
        update_doc(path, sa_rows, hy_rows)


if __name__ == "__main__":
    main()
