#!/usr/bin/env python3
"""Build an editable PowerPoint diagram of the simulation process."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x28, 0x33)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
LINE = RGBColor(0x7F, 0x8C, 0x8D)

BLUE = RGBColor(0x2E, 0x86, 0xAB)
GREEN = RGBColor(0x1D, 0x7A, 0x46)
PURPLE = RGBColor(0x6C, 0x34, 0x8D)
AMBER = RGBColor(0xB7, 0x77, 0x0D)
RED = RGBColor(0xA9, 0x32, 0x26)

BLUE_FILL = RGBColor(0xEA, 0xF4, 0xFA)
GREEN_FILL = RGBColor(0xE8, 0xF8, 0xF0)
PURPLE_FILL = RGBColor(0xF5, 0xEE, 0xF8)
AMBER_FILL = RGBColor(0xFE, 0xF9, 0xE7)
RED_FILL = RGBColor(0xFB, 0xEE, 0xEC)
GREY_FILL = RGBColor(0xF4, 0xF6, 0xF7)


def set_run_font(run, size, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_text(shape, paragraphs, font_size=11, color=INK, align=PP_ALIGN.LEFT, anchor="ctr"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.07)
    tf.margin_right = Inches(0.07)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    tf._txBody.bodyPr.set("anchor", anchor)
    first = True
    for item in paragraphs:
        if isinstance(item, str):
            text, opts = item, {}
        else:
            text, opts = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.clear()
        p.alignment = opts.get("align", align)
        p.space_before = Pt(opts.get("space_before", 0))
        p.space_after = Pt(opts.get("space_after", 1))
        run = p.add_run()
        run.text = text
        set_run_font(
            run,
            opts.get("size", font_size),
            bold=opts.get("bold", False),
            color=opts.get("color", color),
        )


def add_box(slide, left, top, width, height, fill, line=LINE, line_width=0.9):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def add_arrow(slide, left, top, width=0.22, height=0.14, fill=NAVY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_right_arrow(slide, left, top, width=0.22, height=0.14, fill=NAVY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    bg = add_rect(slide, 0, 0, 13.333, 7.5, WHITE)
    bg.shadow.inherit = False

    title = add_rect(slide, 0, 0, 13.333, 0.52, NAVY)
    add_text(
        title,
        [("Simulation study: data-generating process, sampling and analysis", {"size": 20, "bold": True, "color": WHITE})],
        align=PP_ALIGN.LEFT,
        anchor="ctr",
    )
    title.text_frame.margin_left = Inches(0.28)

    subtitle = add_rect(slide, 0.28, 0.58, 12.77, 0.28, GREY_FILL)
    add_text(
        subtitle,
        [(
            "18 scenarios  =  3 overlap levels  ×  3 sample sizes  ×  2 trial designs     ·     1,000 replicates     ·     true log HR = log(0.70)",
            {"size": 11, "color": MUTED},
        )],
        align=PP_ALIGN.CENTER,
    )

    # ---- Column headers ----
    headers = [
        (0.22, BLUE, "1. Data-generating process"),
        (4.58, GREEN, "2. Sampling (each replicate)"),
        (8.94, PURPLE, "3. Analysis methods"),
    ]
    for left, colour, label in headers:
        bar = add_rect(slide, left, 0.96, 4.16, 0.36, colour)
        add_text(bar, [(label, {"size": 13, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)

    # ========================= COLUMN 1: DGP =========================
    x0 = 0.22
    w = 4.16

    box = add_box(slide, x0, 1.42, w, 0.72, BLUE_FILL, BLUE)
    add_text(
        box,
        [
            ("Covariates", {"size": 12, "bold": True, "align": PP_ALIGN.CENTER}),
            ("X1, X2 ~ Normal    ·    X3, X4, X5 ~ Bernoulli", {"size": 10, "align": PP_ALIGN.CENTER}),
        ],
        align=PP_ALIGN.CENTER,
    )
    add_arrow(slide, 2.19, 2.16, fill=BLUE)

    box = add_box(slide, x0, 2.32, w, 1.42, GREY_FILL, MUTED)
    add_text(
        box,
        [
            ("Four super-populations (n = 5,000 each)", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "space_after": 4}),
            ("External control: untreated; large-overlap margins", {"size": 10}),
            ("Trial (no overlap): shifted means / probabilities", {"size": 10}),
            ("Trial (moderate): halfway covariate shift", {"size": 10}),
            ("Trial (large overlap): same margins as control", {"size": 10}),
        ],
        font_size=10,
    )
    add_arrow(slide, 2.19, 3.76, fill=BLUE)

    box = add_box(slide, x0, 3.92, w, 1.28, AMBER_FILL, AMBER)
    add_text(
        box,
        [
            ("Exponential proportional hazards", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "space_after": 4}),
            ("λ = exp(γ0 + γ1X1 + … + γ5X5 + β·trt)", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("β = log(0.70) in the trial population only", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("Random censoring (rate 0.02); trial also administratively censored at 6 years", {"size": 10, "align": PP_ALIGN.CENTER}),
        ],
        font_size=10,
    )

    # Overlap cards
    add_text_box = add_box(slide, x0, 5.32, 1.32, 1.48, RED_FILL, RED)
    add_text(
        add_text_box,
        [
            ("No overlap", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "color": RED}),
            ("X1 ~ N(1.5, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X2 ~ N(0.5, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X3 ~ Bern(0.60)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X4 ~ Bern(0.70)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X5 ~ Bern(0.80)", {"size": 9, "align": PP_ALIGN.CENTER}),
        ],
        font_size=9,
    )
    mid = add_box(slide, x0 + 1.42, 5.32, 1.32, 1.48, AMBER_FILL, AMBER)
    add_text(
        mid,
        [
            ("Moderate", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "color": AMBER}),
            ("X1 ~ N(0.75, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X2 ~ N(0.25, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X3 ~ Bern(0.45)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X4 ~ Bern(0.55)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X5 ~ Bern(0.65)", {"size": 9, "align": PP_ALIGN.CENTER}),
        ],
        font_size=9,
    )
    good = add_box(slide, x0 + 2.84, 5.32, 1.32, 1.48, GREEN_FILL, GREEN)
    add_text(
        good,
        [
            ("Large overlap", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "color": GREEN}),
            ("X1 ~ N(0, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X2 ~ N(0, 1)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X3 ~ Bern(0.30)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X4 ~ Bern(0.40)", {"size": 9, "align": PP_ALIGN.CENTER}),
            ("X5 ~ Bern(0.50)", {"size": 9, "align": PP_ALIGN.CENTER}),
        ],
        font_size=9,
    )

    # ========================= COLUMN 2: SAMPLING =========================
    x1 = 4.58

    box = add_box(slide, x1, 1.42, w, 1.18, GREEN_FILL, GREEN)
    add_text(
        box,
        [
            ("From the control super-population", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER}),
            ("Sample n_cont historical controls", {"size": 10, "align": PP_ALIGN.CENTER, "space_before": 3}),
            ("Small 300   ·   Medium 400   ·   Large 500", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("without replacement", {"size": 9, "align": PP_ALIGN.CENTER, "color": MUTED}),
        ],
        font_size=10,
    )
    add_arrow(slide, 6.55, 2.62, fill=GREEN)

    box = add_box(slide, x1, 2.78, w, 1.18, GREEN_FILL, GREEN)
    add_text(
        box,
        [
            ("From the chosen trial super-population", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER}),
            ("Sample n_trt prospective patients", {"size": 10, "align": PP_ALIGN.CENTER, "space_before": 3}),
            ("Small 50   ·   Medium 75   ·   Large 100", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("without replacement", {"size": 9, "align": PP_ALIGN.CENTER, "color": MUTED}),
        ],
        font_size=10,
    )
    add_arrow(slide, 6.55, 3.98, fill=GREEN)

    sa = add_box(slide, x1, 4.16, 2.00, 2.64, BLUE_FILL, BLUE)
    add_text(
        sa,
        [
            ("Fully synthetic", {"size": 12, "bold": True, "align": PP_ALIGN.CENTER, "color": BLUE}),
            ("single-arm design", {"size": 10, "align": PP_ALIGN.CENTER, "color": BLUE, "space_after": 6}),
            ("All n_trt receive", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("experimental treatment", {"size": 10, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Concurrent controls: 0", {"size": 10, "bold": True, "align": PP_ALIGN.CENTER}),
            ("Historical controls: n_cont", {"size": 10, "align": PP_ALIGN.CENTER}),
        ],
        font_size=10,
    )
    hy = add_box(slide, x1 + 2.16, 4.16, 2.00, 2.64, AMBER_FILL, AMBER)
    add_text(
        hy,
        [
            ("Partially synthetic", {"size": 12, "bold": True, "align": PP_ALIGN.CENTER, "color": AMBER}),
            ("hybrid design", {"size": 10, "align": PP_ALIGN.CENTER, "color": AMBER, "space_after": 6}),
            ("n_trt randomised 1:1", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("experimental vs control", {"size": 10, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Concurrent controls: n_trt/2", {"size": 10, "bold": True, "align": PP_ALIGN.CENTER}),
            ("Historical controls: n_cont", {"size": 10, "align": PP_ALIGN.CENTER}),
        ],
        font_size=10,
    )

    # ========================= COLUMN 3: ANALYSES =========================
    x2 = 8.94

    box = add_box(slide, x2, 1.42, w, 0.72, PURPLE_FILL, PURPLE)
    add_text(
        box,
        [
            ("Counterfactual model (CFM)", {"size": 12, "bold": True, "align": PP_ALIGN.CENTER, "color": PURPLE}),
            ("flexsurvspline, 3 knots, fitted to historical controls", {"size": 10, "align": PP_ALIGN.CENTER}),
            ("covariates X1–X5", {"size": 10, "align": PP_ALIGN.CENTER}),
        ],
        font_size=10,
    )
    add_arrow(slide, 10.91, 2.16, fill=PURPLE)

    sa_m = add_box(slide, x2, 2.32, 2.00, 4.48, BLUE_FILL, BLUE)
    add_text(
        sa_m,
        [
            ("Single-arm methods", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "color": BLUE, "space_after": 8}),
            ("Pooled Cox", {"size": 10, "bold": True}),
            ("unadjusted", {"size": 10, "space_after": 8}),
            ("Personalised SC (PSC)", {"size": 10, "bold": True}),
            ("full CFM", {"size": 10, "space_after": 8}),
            ("Population SC", {"size": 10, "bold": True}),
            ("entropy balance", {"size": 10, "space_after": 8}),
            ("Bayesian historical", {"size": 10, "bold": True}),
            ("informative prior", {"size": 10}),
            ("case-weighted", {"size": 10}),
        ],
        font_size=10,
        align=PP_ALIGN.LEFT,
        anchor="t",
    )

    hy_m = add_box(slide, x2 + 2.16, 2.32, 2.00, 4.48, AMBER_FILL, AMBER)
    add_text(
        hy_m,
        [
            ("Hybrid methods", {"size": 11, "bold": True, "align": PP_ALIGN.CENTER, "color": AMBER, "space_after": 8}),
            ("RCT only", {"size": 10, "bold": True}),
            ("unadjusted Cox", {"size": 10, "space_after": 8}),
            ("Pooled Cox", {"size": 10, "bold": True}),
            ("unadjusted", {"size": 10, "space_after": 8}),
            ("PSC (combined)", {"size": 10, "bold": True}),
            ("full CFM", {"size": 10, "space_after": 8}),
            ("Population SC", {"size": 10, "bold": True}),
            ("entropy balance", {"size": 10, "space_after": 8}),
            ("Bayesian borrowing", {"size": 10, "bold": True}),
            ("vague prior", {"size": 10}),
            ("informative prior", {"size": 10}),
            ("commensurate", {"size": 10}),
            ("case-weighted", {"size": 10}),
        ],
        font_size=10,
        align=PP_ALIGN.LEFT,
        anchor="t",
    )

    # Column separators
    for left in (4.42, 8.78):
        line = add_rect(slide, left, 0.96, 0.015, 5.84, RGBColor(0xD5, 0xD8, 0xDC))

    out = "/Users/richardjackson/Documents/GitHub/synthetic-Controlled-Trials/Manuscript/Supplementary_Figure_simulation_process.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
